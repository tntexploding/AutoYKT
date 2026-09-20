"""Extract text from common courseware files and ingest it locally."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

from autoykt.knowledge.store import IngestResult, KnowledgeStore


_TEXT_SUFFIXES = frozenset({".txt", ".md", ".rst", ".csv"})
_SUPPORTED_SUFFIXES = _TEXT_SUFFIXES | {".pdf", ".pptx", ".docx"}


class UnsupportedCoursewareError(ValueError):
    """Raised when a courseware format has no configured extractor."""


@dataclass(frozen=True)
class FileIngestResult:
    """One file ingestion outcome."""

    path: Path
    result: IngestResult


def collect_courseware(paths: list[Path]) -> list[Path]:
    """Expand files and directories into supported courseware paths."""
    files: set[Path] = set()
    for path in paths:
        resolved = path.expanduser().resolve()
        if resolved.is_dir():
            files.update(
                item
                for item in resolved.rglob("*")
                if item.is_file() and item.suffix.lower() in _SUPPORTED_SUFFIXES
            )
        elif resolved.is_file():
            files.add(resolved)
        else:
            raise FileNotFoundError(f"courseware path not found: {resolved}")
    return sorted(files)


def extract_courseware_text(path: Path) -> str:
    """Extract readable text from a supported courseware file."""
    suffix = path.suffix.lower()
    if suffix in _TEXT_SUFFIXES:
        return path.read_text(encoding="utf-8-sig")
    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix == ".pptx":
        return _extract_powerpoint(path)
    if suffix == ".docx":
        return _extract_word(path)
    raise UnsupportedCoursewareError(
        f"unsupported courseware format '{suffix}': {path}"
    )


def ingest_courseware(
    store: KnowledgeStore,
    course_id: str,
    paths: list[Path],
) -> list[FileIngestResult]:
    """Extract and add every supplied courseware file."""
    results: list[FileIngestResult] = []
    for path in collect_courseware(paths):
        text = extract_courseware_text(path)
        result = store.add_document(
            course_id=course_id,
            source=str(path),
            title=path.stem,
            text=text,
        )
        results.append(FileIngestResult(path=path, result=result))
    return results


def _extract_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader  # pylint: disable=import-outside-toplevel
    except ImportError as error:
        raise RuntimeError(
            "PDF ingestion requires the 'pypdf' package"
        ) from error
    reader = PdfReader(str(path))
    return "\n\n".join(page.extract_text() or "" for page in reader.pages)


def _extract_powerpoint(path: Path) -> str:
    try:
        from pptx import (  # pylint: disable=import-outside-toplevel
            Presentation,
        )
    except ImportError as error:
        raise RuntimeError(
            "PowerPoint ingestion requires the 'python-pptx' package"
        ) from error
    presentation = Presentation(str(path))
    slides: list[str] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        texts = []
        for shape in slide.shapes:
            texts.extend(_powerpoint_shape_text(shape))
        if texts:
            slides.append(f"第 {slide_number} 页\n" + "\n".join(texts))
    return "\n\n".join(slides)


def _powerpoint_shape_text(shape: Any) -> list[str]:
    texts = []
    text = str(getattr(shape, "text", "")).strip()
    if text:
        texts.append(text)
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            texts.append(" | ".join(cell.text for cell in row.cells))
    for child in getattr(shape, "shapes", ()):
        texts.extend(_powerpoint_shape_text(child))
    return texts


def _extract_word(path: Path) -> str:
    try:
        from docx import Document  # pylint: disable=import-outside-toplevel
    except ImportError as error:
        raise RuntimeError(
            "Word ingestion requires the 'python-docx' package"
        ) from error
    document = Document(str(path))
    texts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            texts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(texts)
