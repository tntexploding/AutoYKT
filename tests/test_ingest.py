"""Tests for local courseware extraction adapters."""

from pathlib import Path
import tempfile
import unittest

from docx import Document
from pptx import Presentation
from pptx.util import Inches

from autoykt.knowledge.ingest import (
    collect_courseware,
    extract_courseware_text,
)


class CoursewareIngestTest(unittest.TestCase):
    """Exercise text, Word, PowerPoint, and directory discovery."""

    def test_extracts_text_and_ignores_unsupported_directory_files(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            text_path = root / "lesson.md"
            text_path.write_text("课程知识", encoding="utf-8")
            (root / "ignored.bin").write_bytes(b"ignored")

            self.assertEqual(collect_courseware([root]), [text_path.resolve()])
            self.assertEqual(extract_courseware_text(text_path), "课程知识")

    def test_extracts_word_paragraphs(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "lesson.docx"
            document = Document()
            document.add_paragraph("牛顿第二定律")
            document.save(str(path))

            self.assertIn("牛顿第二定律", extract_courseware_text(path))

    def test_extracts_powerpoint_shape_text_with_slide_number(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "lesson.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            box = slide.shapes.add_textbox(
                Inches(1), Inches(1), Inches(4), Inches(1)
            )
            box.text = "动量守恒"
            presentation.save(str(path))

            extracted = extract_courseware_text(path)
            self.assertIn("第 1 页", extracted)
            self.assertIn("动量守恒", extracted)


class CoursewareTableTest(unittest.TestCase):
    """Course knowledge often lives in tables or grouped slide shapes."""

    def test_word_tables_are_ingested(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "table.docx"
            document = Document()
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "课程定义"
            table.cell(0, 1).text = "专有知识内容"
            document.save(str(path))
            self.assertIn("专有知识内容", extract_courseware_text(path))

    def test_powerpoint_tables_and_groups_are_ingested(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "table.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            table = slide.shapes.add_table(
                1, 1, Inches(0), Inches(0), Inches(4), Inches(1)
            ).table
            table.cell(0, 0).text = "课程表格内容"
            group = slide.shapes.add_group_shape()
            group.shapes.add_textbox(
                Inches(0), Inches(0), Inches(2), Inches(1)
            ).text = "组合中的内容"
            presentation.save(str(path))
            extracted = extract_courseware_text(path)
            self.assertIn("课程表格内容", extracted)
            self.assertIn("组合中的内容", extracted)


if __name__ == "__main__":
    unittest.main()
